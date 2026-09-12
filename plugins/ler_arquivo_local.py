"""Plugin de leitura de arquivos locais pro chattui.

Lê .txt/.md/.csv/.log direto e PDF/DOCX extraindo o texto (pypdf /
python-docx — instalados por requirements-rag-opcional.txt). Complementa
o RAG (que só indexa .txt/.md): no chat, o modelo pode ler um PDF e
resumir, ou extrair um trecho específico — e o resultado vira nota
indexável automaticamente.

Fica em plugins/ (pasta do projeto); pra desativar, renomeie começando
com "_".
"""

from __future__ import annotations

from pathlib import Path

MAX_CHARS = 6000

TOOL_SCHEMA = {
    "name": "ler_arquivo_local",
    "description": (
        "Lê o texto de um arquivo local do computador do usuário: .txt, .md, .csv, "
        ".log, .json, .pdf ou .docx. Use quando ele pedir pra resumir/consultar um "
        "arquivo pelo caminho (ex: 'resume o PDF em ~/Downloads/x.pdf'). "
        "Caminhos com ~ são aceitos."
    ),
    "parameters": {
        "type": "object",
        "properties": {
            "caminho": {"type": "string", "description": "caminho do arquivo (ex: ~/Downloads/edital.pdf)"},
            "max_chars": {"type": "integer", "description": f"limite de caracteres retornados (padrão {MAX_CHARS})"},
        },
        "required": ["caminho"],
    },
}

_PLAIN_EXTS = {".txt", ".md", ".csv", ".log", ".json", ".py", ".toml", ".html", ".xml"}


_BLOCKED_DIRS = {".ssh", ".gnupg", ".aws", ".kube", ".docker", "credentials", "secrets"}
_BLOCKED_SUFFIXES = (".pem", ".key", ".p12", ".pfx", ".kdbx")
_BLOCKED_NAME_PATTERNS = ("id_", ".env", "token", "secret", "credential", "senha", "password")


def _caminho_sensivel(path: Path) -> str | None:
    """Deny-list de caminhos sensíveis (evita exfiltração de segredos)."""
    parts = path.parts
    for part in parts[:-1]:
        if part.startswith(".") or part.lower() in _BLOCKED_DIRS:
            return f"diretório oculto/sensível: {part!r}"
    name = path.name.lower()
    if name.startswith(".") :
        return "arquivo oculto (dotfile)"
    if name.endswith(_BLOCKED_SUFFIXES):
        return "extensão típica de chave/certificado"
    for pattern in _BLOCKED_NAME_PATTERNS:
        if pattern in name:
            return f"padrão sensível no nome: {pattern!r}"
    return None


def _extract(path: Path, budget: int) -> str:
    ext = path.suffix.lower()
    if ext in _PLAIN_EXTS:
        text = path.read_text(errors="replace")
        return text if len(text) <= budget else text[:budget] + f"\n\n[...conteúdo truncado em {budget} caracteres...]"

    if ext == ".pdf":
        try:
            from pypdf import PdfReader  # type: ignore[import-not-found]
        except ImportError:
            return "[erro] pypdf não instalado — rode: pip install -r requirements-rag-opcional.txt"
        reader = PdfReader(str(path))
        parts: list[str] = []
        total = 0
        for page in reader.pages:
            page_text = page.extract_text() or ""
            parts.append(page_text)
            total += len(page_text)
            if total >= budget:
                break
        text = "\n\n".join(parts)
        n_pages = len(reader.pages)
        tail = f"\n\n[...trecho de {n_pages} página(s), truncado em {budget} caracteres...]" if total > budget else ""
        return text[:budget] + tail

    if ext == ".docx":
        try:
            from docx import Document  # type: ignore[import-not-found]
        except ImportError:
            return "[erro] python-docx não instalado — rode: pip install -r requirements-rag-opcional.txt"
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
        return text if len(text) <= budget else text[:budget] + f"\n\n[...conteúdo truncado em {budget} caracteres...]"

    return (
        f"[erro] extensão {ext or '(sem extensão)'} não suportada. "
        "Suporto: .txt .md .csv .log .json .pdf .docx"
    )


def run(caminho: str, max_chars: int = MAX_CHARS) -> str:
    if not caminho or not caminho.strip():
        return "[erro] informe o caminho do arquivo."
    path = Path(caminho.strip()).expanduser()
    bloqueio = _caminho_sensivel(path)
    if bloqueio:
        return (
            f"[erro] caminho sensível bloqueado ({bloqueio}). "
            "Deny-list do plugin — edite ler_arquivo_local.py se precisar ler esse arquivo."
        )
    if not path.exists():
        return f"[erro] arquivo não encontrado: {path}"
    if path.is_dir():
        return f"[erro] é uma pasta, não um arquivo: {path}"
    budget = max(500, int(max_chars or MAX_CHARS))
    try:
        return _extract(path, budget)
    except Exception as exc:  # noqa: BLE001 — arquivo corrompido/pdf criptografado etc.
        return f"[erro ao ler {path.name}] {exc}"
